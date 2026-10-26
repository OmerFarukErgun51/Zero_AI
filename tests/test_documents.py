import openpyxl
from docx import Document
from pptx import Presentation

from zero.documents.excel import make_excel
from zero.documents.powerpoint import make_powerpoint, Slide
from zero.documents.word import make_word, Section


def test_excel_created_with_headers(tmp_path):
    p = make_excel(tmp_path / "a.xlsx", "Rapor", ["Ad", "Sayı"], [["Ali", 3], ["Veli", 5]])
    wb = openpyxl.load_workbook(p)
    ws = wb.active
    assert [c.value for c in ws[1]] == ["Ad", "Sayı"]
    assert ws["A2"].value == "Ali"


def test_pptx_slide_count_and_title(tmp_path):
    slides = [Slide("Giriş", ["madde 1", "madde 2"]), Slide("Sonuç", ["bitti"])]
    p = make_powerpoint(tmp_path / "a.pptx", "Sunum", slides)
    prs = Presentation(p)
    assert len(prs.slides) == len(slides) + 1  # başlık slaytı + içerik
    assert prs.slides[1].shapes.title.text == "Giriş"


def test_docx_has_heading_and_body(tmp_path):
    secs = [Section("Özet", "Bu bir özettir.")]
    p = make_word(tmp_path / "a.docx", "Belge", secs)
    doc = Document(p)
    texts = [para.text for para in doc.paragraphs]
    assert "Özet" in texts
    assert "Bu bir özettir." in texts
