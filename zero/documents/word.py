"""Word (.docx) üretimi — python-docx ile, tamamen offline."""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from docx import Document


@dataclass
class Section:
    heading: str
    body: str


def make_word(path: Path, title: str, sections: list[Section]) -> Path:
    path = Path(path)
    doc = Document()
    doc.add_heading(title, level=0)

    for section in sections:
        doc.add_heading(section.heading, level=1)
        doc.add_paragraph(section.body)

    doc.save(path)
    return path
