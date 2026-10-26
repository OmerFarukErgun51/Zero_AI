"""PowerPoint (.pptx) üretimi — python-pptx ile, tamamen offline."""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from pptx import Presentation
from pptx.util import Pt


@dataclass
class Slide:
    title: str
    bullets: list[str]


def make_powerpoint(path: Path, title: str, slides: list[Slide]) -> Path:
    path = Path(path)
    prs = Presentation()

    # Başlık slaytı
    title_layout = prs.slide_layouts[0]
    s = prs.slides.add_slide(title_layout)
    s.shapes.title.text = title

    # İçerik slaytları
    content_layout = prs.slide_layouts[1]
    for slide in slides:
        sl = prs.slides.add_slide(content_layout)
        sl.shapes.title.text = slide.title
        body = sl.placeholders[1].text_frame
        body.clear()
        for i, bullet in enumerate(slide.bullets):
            para = body.paragraphs[0] if i == 0 else body.add_paragraph()
            para.text = bullet
            para.font.size = Pt(18)

    prs.save(path)
    return path
